import os
import re
import docx
import pdfplumber
import openpyxl

from core.bip39_loader import load_bip39_words
from core.utils import is_wallet_file

# 掃描器主邏輯
class MnemonicScanner:
    def __init__(self, bip39_path):
        self.bip39_words = load_bip39_words(bip39_path)
        self.bip39_set = set(self.bip39_words)
        self.target_exts = {'.txt', '.docx', '.pdf', '.xlsx'}
        self.exclude_files = {
            os.path.abspath('output/key.txt'),
            os.path.abspath('resources/bip39_english.txt')
        }

    def extract_text(self, filepath):
        ext = os.path.splitext(filepath)[1].lower()
        try:
            if ext == ".txt":
                with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()
            elif ext == ".docx":
                doc = docx.Document(filepath)
                return "\n".join(p.text for p in doc.paragraphs)
            elif ext == ".pdf":
                with pdfplumber.open(filepath) as pdf:
                    return "\n".join(page.extract_text() or "" for page in pdf.pages)
            elif ext == ".xlsx":
                wb = openpyxl.load_workbook(filepath, data_only=True)
                text = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value:
                                text.append(str(cell.value))
                return "\n".join(text)
        except Exception:
            return ""
        return ""

    def find_consecutive_mnemonics(self, words, target_lengths={12, 15, 18, 21, 24}):
        count = 0
        for word in words:
            if word in self.bip39_set:
                count += 1
            else:
                if count in target_lengths:
                    return count
                count = 0
        if count in target_lengths:
            return count
        return 0

    def scan_file(self, filepath):
        if os.path.abspath(filepath) in self.exclude_files:
            return None

        if os.path.splitext(filepath)[1].lower() not in self.target_exts:
            return None

        text = self.extract_text(filepath).lower()
        words = re.findall(r'\b[a-z]+\b', text)
        mnemonic_count = self.find_consecutive_mnemonics(words)

        result = {
            "path": filepath,
            "mnemonic_count": mnemonic_count,
            "is_wallet": is_wallet_file(filepath, text)
        }

        if result["mnemonic_count"] > 0 or result["is_wallet"]:
            return result
        return None

    def scan_directory(self, base_path, callback=None):
        results = []
        for root, _, files in os.walk(base_path):
            for file in files:
                filepath = os.path.join(root, file)
                res = self.scan_file(filepath)
                if res:
                    results.append(res)
                if callback:
                    callback(filepath)
        return results
